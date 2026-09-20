from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import fitz
from docx import Document

from rag.weijing.models import ConcernCode, RegionCode, RegionReading, ScoreReading, WeijingReport


MAX_REPORT_BYTES = 10 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}

REGION_ALIASES = {
    RegionCode.F: ("额部", "额头"),
    RegionCode.C: ("下巴", "下颌"),
    RegionCode.L: ("左面部", "左脸", "左侧面部", "左侧面颊", "左颊"),
    RegionCode.R: ("右面部", "右脸", "右侧面部", "右侧面颊", "右颊"),
}
CONCERN_ALIASES = {
    ConcernCode.O: ("油脂评分", "油脂", "油光", "出油"),
    ConcernCode.P: ("综合色素", "色斑评分", "色斑", "色素"),
    ConcernCode.D: ("干燥性细纹", "干燥细纹", "细纹评分", "细纹"),
}
LEVEL_LABELS = ("未见明显", "轻度", "中度", "较明显", "显著")
_LEVEL_PATTERN = "|".join(LEVEL_LABELS)


MAX_REPORT_BYTES = 10 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}

REGION_ALIASES = {
    RegionCode.F: ("额部", "额头"),
    RegionCode.C: ("下巴", "下颌"),
    RegionCode.L: ("左面部", "左脸", "左侧面部"),
    RegionCode.R: ("右面部", "右脸", "右侧面部"),
}
CONCERN_ALIASES = {
    ConcernCode.O: ("油脂", "油脂评分", "油光"),
    ConcernCode.P: ("色斑", "色斑评分", "综合色素"),
    ConcernCode.D: ("干燥性细纹", "干燥细纹", "细纹"),
}
LEVEL_LABELS = ("未见明显", "轻度", "中度", "较明显", "显著")


class ReportUploadError(ValueError):
    pass


class ReportParseError(ValueError):
    pass


def parse_weijing_report(data: bytes, filename: str) -> WeijingReport:
    extension = _validate_upload(data, filename)
    text = _extract_text(data, extension)
    normalized = re.sub(r"[\t\u3000]+", " ", text.replace("\r", "\n"))
    regions = _merge_region_maps(
        _parse_json_regions(normalized),
        _parse_table_regions(normalized),
        _parse_labeled_regions(normalized),
    )
    notes: list[str] = []
    for region, aliases in REGION_ALIASES.items():
        if region not in regions:
            notes.append(f"缺少{aliases[0]}数据")

    if not regions or not any(
        reading.score is not None
        for region in regions.values()
        for reading in region.scores.values()
    ):
        raise ReportParseError(_parse_failure_message(normalized))

    excluded = [name for name in ("鼻部", "鼻子", "眼周") if name in normalized]
    report_id = _first_value(normalized, ("报告编号", "报告ID", "report_id"))
    algorithm_version = _first_value(normalized, ("算法版本", "algorithm_version"))
    assessed_at = _parse_datetime(
        _first_value(normalized, ("检测时间", "评估时间", "assessed_at"))
    )
    file_sha256 = hashlib.sha256(data).hexdigest()
    return WeijingReport(
        report_id=report_id or f"report-{file_sha256[:24]}",
        assessed_at=assessed_at,
        algorithm_version=algorithm_version or "未提供",
        filename=filename,
        file_sha256=file_sha256,
        regions=regions,
        excluded_regions=list(dict.fromkeys(excluded)),
        parsing_notes=notes,
    )


def _parse_labeled_regions(text: str) -> dict[RegionCode, RegionReading]:
    regions: dict[RegionCode, RegionReading] = {}
    for region, aliases in REGION_ALIASES.items():
        section = _region_section(text, aliases)
        scores = {
            concern: _score_reading(section, aliases, concern)
            for concern in ConcernCode
        }
        if any(
            item.score is not None or item.issue in {"invalid", "conflict"}
            for item in scores.values()
        ):
            regions[region] = RegionReading(scores=scores)
    return regions


def _parse_table_regions(text: str) -> dict[RegionCode, RegionReading]:
    header = re.search(
        rf"(油脂|油光).{{0,80}}(色斑|色素).{{0,80}}(干燥性细纹|细纹)",
        text,
        re.DOTALL,
    )
    if header is None:
        return {}
    regions: dict[RegionCode, dict[ConcernCode, ScoreReading]] = {}
    for region, aliases in REGION_ALIASES.items():
        alias_pattern = "|".join(re.escape(item) for item in aliases)
        match = re.search(
            rf"(?:{alias_pattern})\D{{0,12}}"
            rf"(-?\d+(?:\.\d+)?)\s*(?:分)?\s*({_LEVEL_PATTERN})?\s*"
            rf"(-?\d+(?:\.\d+)?)\s*(?:分)?\s*({_LEVEL_PATTERN})?\s*"
            rf"(-?\d+(?:\.\d+)?)\s*(?:分)?\s*({_LEVEL_PATTERN})?",
            text,
        )
        if match is None:
            continue
        regions[region] = {
            ConcernCode.O: _parse_score_match(match, 1, 2),
            ConcernCode.P: _parse_score_match(match, 3, 4),
            ConcernCode.D: _parse_score_match(match, 5, 6),
        }
    return {
        region: RegionReading(scores=scores)
        for region, scores in regions.items()
    }


def _parse_json_regions(text: str) -> dict[RegionCode, RegionReading]:
    regions: dict[RegionCode, RegionReading] = {}
    for match in re.finditer(r"\{[^{}]{20,8000}\}", text):
        try:
            payload = json.loads(match.group())
        except json.JSONDecodeError:
            continue
        raw_regions = payload.get("regions")
        if not isinstance(raw_regions, dict):
            continue
        for key, value in raw_regions.items():
            try:
                region = RegionCode(str(key))
            except ValueError:
                continue
            if not isinstance(value, dict):
                continue
            scores = {}
            source = value.get("scores") if isinstance(value.get("scores"), dict) else value
            for concern in ConcernCode:
                item = source.get(concern.value) if isinstance(source, dict) else None
                scores[concern] = _score_from_json(item)
            if any(item.score is not None or item.issue for item in scores.values()):
                regions[region] = RegionReading(scores=scores)
    return regions


def _score_from_json(value) -> ScoreReading:
    if value is None:
        return ScoreReading(issue="missing")
    if isinstance(value, bool) or isinstance(value, str) and not value.isdigit():
        if isinstance(value, dict):
            return _score_from_json(value.get("score"))
        return ScoreReading(issue="invalid")
    if isinstance(value, dict):
        raw = value.get("score")
        label = value.get("label") or value.get("reported_label")
        reading = _score_from_json(raw)
        if label and reading.score is not None:
            return reading.model_copy(update={"reported_label": str(label)})
        return reading
    if isinstance(value, float) and not value.is_integer():
        return ScoreReading(issue="invalid")
    try:
        score = int(value)
    except (TypeError, ValueError):
        return ScoreReading(issue="invalid")
    if not 0 <= score <= 100:
        return ScoreReading(issue="invalid")
    return ScoreReading(score=score)


def _merge_region_maps(*maps: dict[RegionCode, RegionReading]) -> dict[RegionCode, RegionReading]:
    merged: dict[RegionCode, dict[ConcernCode, ScoreReading]] = {}
    for mapping in maps:
        for region, reading in mapping.items():
            current = merged.setdefault(region, {})
            for concern, score in reading.scores.items():
                previous = current.get(concern)
                if previous is None or previous.score is None:
                    current[concern] = score
    result = {}
    for region, scores in merged.items():
        filled = {
            concern: scores.get(concern, ScoreReading(issue="missing"))
            for concern in ConcernCode
        }
        result[region] = RegionReading(scores=filled)
    return result


def _parse_failure_message(text: str) -> str:
    if "泛红" in text or "AISIA_CLINICAL_FACTS" in text:
        return (
            "当前文件是泛红检测报告，未包含四区油脂、色斑、干燥性细纹三项评分，"
            "无法生成微晶改善方案。"
        )
    return "检测报告中未找到四区油脂、色斑或干燥性细纹评分"


def _score_reading(text: str, region_aliases: tuple[str, ...], concern: ConcernCode) -> ScoreReading:
    concern_pattern = "|".join(re.escape(value) for value in CONCERN_ALIASES[concern])
    region_pattern = "|".join(re.escape(value) for value in region_aliases)
    boolean = re.search(
        rf"(?:{concern_pattern})\s*(?:评分|得分)?\s*[:：]?\s*(true|false|True|False|是|否)\b",
        text,
    )
    if boolean:
        return ScoreReading(issue="invalid")
    pattern = (
        rf"(?:{concern_pattern})\s*(?:评分|得分)?\s*[:：]?\s*(-?\d+(?:\.\d+)?)\s*(?:分)?\s*"
        rf"({_LEVEL_PATTERN})?"
    )
    matches = list(re.finditer(pattern, text, re.DOTALL))
    if not matches:
        fallback = re.search(
            rf"(?:{region_pattern}).{{0,40}}?(?:{concern_pattern})\s*[:：]?\s*(-?\d+(?:\.\d+)?)\s*(?:分)?\s*"
            rf"({_LEVEL_PATTERN})?",
            text,
            re.DOTALL,
        )
        matches = [fallback] if fallback else []
    if not matches:
        return ScoreReading(issue="missing")
    parsed = [_parse_score_match(match) for match in matches]
    scores = {item.score for item in parsed if item.score is not None}
    if any(item.issue == "invalid" for item in parsed):
        return next(item for item in parsed if item.issue == "invalid")
    if len(scores) > 1 or any(item.issue == "conflict" for item in parsed):
        first = parsed[0]
        return ScoreReading(score=first.score, reported_label=first.reported_label, issue="conflict")
    return parsed[0]


def _parse_score_match(match: re.Match[str], score_group: int = 1, label_group: int = 2) -> ScoreReading:
    raw = match.group(score_group)
    reported_label = match.group(label_group) or None
    if "." in raw:
        return ScoreReading(reported_label=reported_label, issue="invalid")
    score = int(raw)
    if not 0 <= score <= 100:
        return ScoreReading(reported_label=reported_label, issue="invalid")
    expected = _level_label(score)
    if reported_label and reported_label != expected:
        return ScoreReading(score=score, reported_label=reported_label, issue="conflict")
    return ScoreReading(score=score, reported_label=reported_label)


def _region_section(text: str, aliases: tuple[str, ...]) -> str:
    starts = [text.find(alias) for alias in aliases if text.find(alias) >= 0]
    if not starts:
        return ""
    start = min(starts)
    ends = []
    for other_aliases in REGION_ALIASES.values():
        if other_aliases == aliases:
            continue
        for alias in other_aliases:
            position = text.find(alias, start + 1)
            if position >= 0:
                ends.append(position)
    return text[start:min(ends) if ends else len(text)]


def _level_label(score: int) -> str:
    if score >= 81:
        return "未见明显"
    if score >= 61:
        return "轻度"
    if score >= 41:
        return "中度"
    if score >= 21:
        return "较明显"
    return "显著"


def _first_value(text: str, labels: tuple[str, ...]) -> str | None:
    for label in labels:
        match = re.search(rf"{re.escape(label)}\s*[:：]\s*([^\n|]+)", text, re.I)
        if match:
            return match.group(1).strip()
    return None


def _parse_datetime(value: str | None) -> datetime:
    if value:
        normalized = value.replace("年", "-").replace("月", "-").replace("日", " ").strip()
        try:
            parsed = datetime.fromisoformat(normalized)
            return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)
        except ValueError:
            pass
    return datetime.now(timezone.utc)


def _validate_upload(data: bytes, filename: str) -> str:
    name = (filename or "").strip()
    if not name or Path(name).name != name or "/" in name or "\\" in name:
        raise ReportUploadError("检测报告文件名无效")
    extension = Path(name).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ReportUploadError("仅支持 PDF 或 DOCX 检测报告")
    if not data:
        raise ReportUploadError("检测报告文件为空")
    if len(data) > MAX_REPORT_BYTES:
        raise ReportUploadError("检测报告文件不能超过 10 MB")
    if extension == ".pdf" and not data.startswith(b"%PDF-"):
        raise ReportUploadError("文件扩展名与 PDF 内容不一致")
    if extension == ".docx" and not data.startswith(b"PK"):
        raise ReportUploadError("文件扩展名与 DOCX 内容不一致")
    return extension


def _extract_text(data: bytes, extension: str) -> str:
    try:
        if extension == ".pdf":
            with fitz.open(stream=data, filetype="pdf") as document:
                return "\n".join(page.get_text("text") for page in document)
        document = Document(BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except Exception as exc:
        raise ReportParseError(f"无法读取检测报告：{exc}") from exc
